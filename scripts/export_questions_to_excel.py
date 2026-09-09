"""
ARUNACHAL EXAM PREP - SUPABASE TO EXCEL EXPORT TOOL
===================================================
Exports questions from Supabase into clean, pre-formatted Excel (.xlsx) files.
Enables rapid visual inspection, editing, and offline curation.
"""

import os
import sys
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import requests

SUPABASE_URL = os.environ.get('SUPABASE_URL', 'https://fllopztywwblbucvaths.supabase.co')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY', 'sb_publishable_y68QKKHxBTZxBP3Sf1X7tw_zfFnXX8M')

HEADERS = {
    'apikey': SUPABASE_KEY,
    'Authorization': f'Bearer {SUPABASE_KEY}',
    'Content-Type': 'application/json',
}

def clean_opt(options, idx):
    if not isinstance(options, list) or idx >= len(options):
        return "", ""
    item = options[idx]
    if isinstance(item, dict):
        text = str(item.get('text', '')).strip()
        img = str(item.get('image', item.get('image_url', ''))).strip()
        # Strip leading (a), (b), etc.
        import re
        text = re.sub(r'^\s*[\(\[\{]?[a-dA-D][\)\]\}\.\s]\s*', '', text)
        return text, img
    elif isinstance(item, str):
        import re
        text = re.sub(r'^\s*[\(\[\{]?[a-dA-D][\)\]\}\.\s]\s*', '', item)
        return text.strip(), ""
    return "", ""

def main():
    parser = argparse.ArgumentParser(description="Export Supabase questions to Excel")
    parser.add_argument("--test-id", help="UUID of specific test")
    parser.add_argument("--exam-code", help="Filter by exam code (e.g. APSSB-CGLE)")
    parser.add_argument("--year", help="Filter by year (e.g. 2024)")
    parser.add_argument("--limit", type=int, default=1000, help="Max questions to export")
    parser.add_argument("--out", default="exported_questions.xlsx", help="Output .xlsx file path")
    args = parser.parse_args()

    # Build query
    query = f"{SUPABASE_URL}/rest/v1/questions?select=*,question_groups(*),tests!inner(*)&order=question_number.asc&limit={args.limit}"
    if args.test_id:
        query += f"&test_id=eq.{args.test_id}"
    if args.exam_code:
        query += f"&tests.exam_code=eq.{args.exam_code}"
    if args.year:
        query += f"&tests.year=eq.{args.year}"

    print(f"Fetching questions from Supabase...")
    res = requests.get(query, headers=HEADERS)
    if res.status_code != 200:
        print(f"Error fetching from Supabase ({res.status_code}): {res.text}")
        sys.exit(1)

    questions = res.json()
    print(f"Found {len(questions)} questions. Formatting into Excel...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questions"

    headers = [
        ("test_code", 18), ("year", 8), ("paper_type", 12), ("subject", 22),
        ("topic", 18), ("question_num", 14), ("group_code", 16), ("group_type", 14),
        ("direction", 35), ("passage_text", 45), ("passage_image", 22),
        ("question_text", 50), ("format", 12), ("question_image", 22),
        ("option_a", 25), ("option_a_image", 18),
        ("option_b", 25), ("option_b_image", 18),
        ("option_c", 25), ("option_c_image", 18),
        ("option_d", 25), ("option_d_image", 18),
        ("correct_answer", 14), ("explanation", 45), ("explanation_image", 22)
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[1].height = 30
    for col_idx, (col_name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_font = Font(name="Segoe UI", size=10)
    row_align = Alignment(vertical="top", wrap_text=True)

    for r_idx, q in enumerate(questions, start=2):
        tests = q.get('tests') or {}
        groups = q.get('question_groups') or {}
        options = q.get('options') or []

        opt_a_text, opt_a_img = clean_opt(options, 0)
        opt_b_text, opt_b_img = clean_opt(options, 1)
        opt_c_text, opt_c_img = clean_opt(options, 2)
        opt_d_text, opt_d_img = clean_opt(options, 3)

        row_vals = [
            tests.get('exam_code', ''),
            tests.get('year', ''),
            tests.get('paper_type', 'PYQ'),
            q.get('subject', 'General Studies'),
            q.get('topic', ''),
            q.get('question_number', r_idx - 1),
            groups.get('title', ''),
            groups.get('group_type', 'comprehension'),
            q.get('direction_text') or groups.get('direction_text') or groups.get('instructions') or '',
            groups.get('passage_text', ''),
            groups.get('passage_image_url') or groups.get('image_url', ''),
            q.get('question_text', ''),
            q.get('content_format', 'text'),
            q.get('question_image_url') or q.get('image_url', ''),
            opt_a_text, opt_a_img,
            opt_b_text, opt_b_img,
            opt_c_text, opt_c_img,
            opt_d_text, opt_d_img,
            q.get('correct_option') or q.get('correct_answer', 'a'),
            q.get('explanation', ''),
            q.get('explanation_image_url', '')
        ]

        ws.row_dimensions[r_idx].height = 45
        fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if r_idx % 2 == 0 else PatternFill(fill_type=None)

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = row_font
            cell.alignment = row_align
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill

    wb.save(args.out)
    print(f"Export complete! Saved {len(questions)} questions to '{args.out}'.")

if __name__ == '__main__':
    main()
